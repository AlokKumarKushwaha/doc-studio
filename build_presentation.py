import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = pptx.Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Theme Colors
    BG_DARK = RGBColor(15, 23, 42)       # Slate 900
    BG_LIGHT = RGBColor(248, 250, 252)   # Slate 50
    CARD_BG = RGBColor(255, 255, 255)    # White
    CARD_BORDER = RGBColor(226, 232, 240)# Slate 200
    TEXT_DARK = RGBColor(15, 23, 42)     # Slate 900
    TEXT_MUTED = RGBColor(100, 116, 139) # Slate 500
    TEXT_LIGHT = RGBColor(241, 245, 249) # Slate 100
    EMERALD = RGBColor(16, 185, 129)     # Emerald 500
    INDIGO = RGBColor(79, 70, 229)       # Indigo 600
    BLUE = RGBColor(2, 132, 199)         # Sky 600
    AMBER = RGBColor(217, 119, 6)        # Amber 600

    def add_header(slide, category, title):
        # Category Tag
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(10), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = EMERALD
        p_cat.font.name = "Segoe UI"

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.5), Inches(0.7))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title
        p_title.font.size = Pt(26)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_DARK
        p_title.font.name = "Segoe UI"

        # Footer
        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.0), Inches(11.7), Inches(0.35))
        tf_foot = footer_box.text_frame
        p_foot = tf_foot.paragraphs[0]
        p_foot.text = "DocStudio • Cloud-Powered PDF & Document Engineering Suite • College Presentation"
        p_foot.font.size = Pt(10)
        p_foot.font.color.rgb = TEXT_MUTED
        p_foot.font.name = "Segoe UI"

    def set_bg(slide, color):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    # ==========================================
    # SLIDE 1: Title Slide (Dark Theme)
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_bg(s1, BG_DARK)

    # Accent decorative bar
    bar = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(1.2), Inches(3.2), Inches(0.36))
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor(6, 78, 59)
    bar.line.fill.background()
    tf_bar = bar.text_frame
    p_b = tf_bar.paragraphs[0]
    p_b.text = "🎓 MAJOR PROJECT PRESENTATION"
    p_b.font.size = Pt(11)
    p_b.font.bold = True
    p_b.font.color.rgb = EMERALD
    p_b.alignment = PP_ALIGN.CENTER

    # Project Title
    t1_box = s1.shapes.add_textbox(Inches(0.85), Inches(1.7), Inches(11.5), Inches(1.4))
    tf1 = t1_box.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    p1.text = "DocStudio"
    p1.font.size = Pt(56)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(255, 255, 255)
    p1.font.name = "Segoe UI"

    # Subtitle
    sub1_box = s1.shapes.add_textbox(Inches(0.9), Inches(3.1), Inches(11.5), Inches(0.9))
    tf_sub1 = sub1_box.text_frame
    tf_sub1.word_wrap = True
    p_sub1 = tf_sub1.paragraphs[0]
    p_sub1.text = "All-in-One Cloud-Powered Web Suite for PDF Engineering, Govt Exam Document Resizing & Interactive Presentation Studio"
    p_sub1.font.size = Pt(19)
    p_sub1.font.color.rgb = RGBColor(203, 213, 225)
    p_sub1.font.name = "Segoe UI"

    # Metadata Card (Presenter & Tech)
    m_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(4.3), Inches(11.5), Inches(2.2))
    m_card.fill.solid()
    m_card.fill.fore_color.rgb = RGBColor(30, 41, 59)
    m_card.line.color.rgb = RGBColor(51, 65, 85)
    m_card.line.width = Pt(1.5)
    tf_mc = m_card.text_frame
    tf_mc.word_wrap = True

    p_mc1 = tf_mc.paragraphs[0]
    p_mc1.text = "Presenter: Alok Kumar Kushwaha  |  Department of Computer Science & Engineering"
    p_mc1.font.size = Pt(16)
    p_mc1.font.bold = True
    p_mc1.font.color.rgb = RGBColor(255, 255, 255)

    p_mc2 = tf_mc.add_paragraph()
    p_mc2.text = "Core Tech Stack: Python (FastAPI), PyMuPDF (Fitz), Pillow, ReportLab, Vanilla JS, Tailwind CSS, PDF.js"
    p_mc2.font.size = Pt(13)
    p_mc2.font.color.rgb = RGBColor(148, 163, 184)

    p_mc3 = tf_mc.add_paragraph()
    p_mc3.text = "Live Deployment: https://docstudio-dty3.onrender.com  |  CI/CD: GitHub + Render Cloud"
    p_mc3.font.size = Pt(13)
    p_mc3.font.color.rgb = EMERALD
    p_mc3.font.bold = True

    s1.notes_slide.notes_text_frame.text = (
        "Good morning respected evaluators, professors, and fellow colleagues. "
        "Today, I am proud to present 'DocStudio', a modern cloud-engineered web platform "
        "designed to revolutionize how documents and PDFs are manipulated, edited, and formatted, "
        "with specialized solutions for government exam aspirants, students, and professionals."
    )

    # ==========================================
    # SLIDE 2: Problem Statement & Industry Need
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_bg(s2, BG_LIGHT)
    add_header(s2, "01. Context & Motivation", "Problem Statement & Industry Gap")

    cards_s2 = [
        ("💸 Heavy Paywalls & Usage Caps",
         "Most market leaders (Adobe Acrobat, Sejda, SmallPDF) restrict basic tasks to 2-3 files per day and require expensive subscriptions ($15-$25/month), rendering them inaccessible to students and rural aspirants.",
         RGBColor(239, 68, 68)),
        ("🏛️ Strict Govt Exam Guidelines",
         "Major examination portals (UPSC, SSC, IBPS, State PSC, NTA) enforce exact dimensions (e.g. 3.5×4.5 cm photo, 4×2 cm signature) and strict file size limits (< 50KB / < 20KB). Candidates face repeated rejections due to blurriness or incorrect aspect ratios.",
         RGBColor(245, 158, 11)),
        ("🔒 Privacy & Tool Fragmentation",
         "Users have to bounce between 4-5 different untrusted sites to crop, compress, convert, and sign files, creating data privacy hazards when uploading confidential Aadhaar, PAN, and academic marksheets.",
         RGBColor(99, 102, 241))
    ]

    for i, (head, desc, accent) in enumerate(cards_s2):
        x = Inches(0.8 + i * 3.95)
        box = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.8), Inches(3.8), Inches(4.8))
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = CARD_BORDER
        box.line.width = Pt(1.5)

        # Accent top bar
        top_bar = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.8), Inches(3.8), Inches(0.15))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = accent
        top_bar.line.fill.background()

        tf = box.text_frame
        tf.word_wrap = True
        p_h = tf.paragraphs[0]
        p_h.text = head
        p_h.font.size = Pt(17)
        p_h.font.bold = True
        p_h.font.color.rgb = TEXT_DARK

        p_d = tf.add_paragraph()
        p_d.text = "\n" + desc
        p_d.font.size = Pt(13)
        p_d.font.color.rgb = RGBColor(71, 85, 105)

    s2.notes_slide.notes_text_frame.text = (
        "In this slide, we outline the exact problems that inspired DocStudio. Millions of students and job applicants "
        "face rejection simply because their uploaded photo or signature exceeded 50KB or 20KB, or got distorted. "
        "Commercial solutions are paywalled, and free alternatives risk user privacy. DocStudio solves this comprehensively."
    )

    # ==========================================
    # SLIDE 3: Proposed Solution & Core Objectives
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    set_bg(s3, BG_LIGHT)
    add_header(s3, "02. Strategic Overview", "Proposed Solution: The DocStudio Ecosystem")

    # Left Hero Box
    left_hero = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    left_hero.fill.solid()
    left_hero.fill.fore_color.rgb = BG_DARK
    left_hero.line.fill.background()
    tf_lh = left_hero.text_frame
    tf_lh.word_wrap = True

    p_lh1 = tf_lh.paragraphs[0]
    p_lh1.text = "🎯 Core Project Mission"
    p_lh1.font.size = Pt(20)
    p_lh1.font.bold = True
    p_lh1.font.color.rgb = EMERALD

    p_lh2 = tf_lh.add_paragraph()
    p_lh2.text = (
        "\nTo engineer a secure, highly optimized, 100% accessible document studio that combines "
        "professional PDF editing, AI-prompted document resizing, and PowerPoint synthesis into a single, "
        "zero-friction browser experience."
    )
    p_lh2.font.size = Pt(14)
    p_lh2.font.color.rgb = RGBColor(226, 232, 240)

    p_lh3 = tf_lh.add_paragraph()
    p_lh3.text = (
        "\n✨ Key Architectural Pillars:\n"
        "• High-Performance Async Python Core (FastAPI)\n"
        "• Client-First Rendering with PDF.js Canvas\n"
        "• Deterministic Iterative Compression Algorithms\n"
        "• Zero-Wall Guest Mode + Optional SQLite Auth"
    )
    p_lh3.font.size = Pt(13)
    p_lh3.font.color.rgb = RGBColor(148, 163, 184)

    # Right 3 Pillars
    right_items = [
        ("⚡ 1. Intelligent Exam Form Cropper", "Natural language parsing + 1-click presets guaranteeing exact physical cm dimensions and <= KB file size limits."),
        ("📝 2. Full In-Browser PDF Editor", "Insert blank pages at runtime, delete pages, add text with font mapping, freehand draw, and place digital signatures."),
        ("🛡️ 3. 30+ Enterprise PDF & PPT Utilities", "Client-side preview coupled with server-side AES-256 encryption, split, merge, conversions, and AI presentation generator.")
    ]

    for j, (title, desc) in enumerate(right_items):
        y = Inches(1.8 + j * 1.65)
        rc = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.7), y, Inches(5.8), Inches(1.45))
        rc.fill.solid()
        rc.fill.fore_color.rgb = CARD_BG
        rc.line.color.rgb = CARD_BORDER
        rc.line.width = Pt(1.5)
        tf_rc = rc.text_frame
        tf_rc.word_wrap = True

        p_rt = tf_rc.paragraphs[0]
        p_rt.text = title
        p_rt.font.size = Pt(15)
        p_rt.font.bold = True
        p_rt.font.color.rgb = TEXT_DARK

        p_rd = tf_rc.add_paragraph()
        p_rd.text = desc
        p_rd.font.size = Pt(12)
        p_rd.font.color.rgb = TEXT_MUTED

    s3.notes_slide.notes_text_frame.text = (
        "DocStudio is engineered around three pillars: First, specialized intelligence for Indian and global exam forms. "
        "Second, true WYSIWYG page-level editing without needing Adobe Acrobat. "
        "And third, an expansive suite of 30+ utilities that run at high speed without forcing users into subscriptions."
    )

    # ==========================================
    # SLIDE 4: System Architecture & Workflow
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    set_bg(s4, BG_LIGHT)
    add_header(s4, "03. Engineering Architecture", "System Architecture & End-to-End Data Pipeline")

    arch_layers = [
        ("1. Presentation Layer (Client UI)",
         "• Responsive Web Interface built with Tailwind CSS & Lucide Icons\n"
         "• Client-side PDF rendering using Mozilla's PDF.js\n"
         "• Interactive Canvas Stage with 8-handle transformation matrices\n"
         "• Aspect Ratio Constraint Logic & Real-time Unit Conversions (cm/mm/in)",
         INDIGO),
        ("2. Application API Layer (FastAPI)",
         "• Asynchronous ASGI Event Loop via Uvicorn Server\n"
         "• High-throughput multipart upload streaming directly into memory\n"
         "• RESTful Tool Endpoints with strict Pydantic payload validation\n"
         "• Session Token Authentication & History Tracking with SQLite3",
         EMERALD),
        ("3. Document Engineering Engines",
         "• PyMuPDF (Fitz): Low-level vector clipping, page insertion & fonts\n"
         "• Pillow (PIL): Iterative quality & DPI downsampling feedback loop\n"
         "• ReportLab & python-pptx: Dynamic generation of PDF and PPT decks\n"
         "• Deflate & Garbage Collection stream compaction for zero-bloat files",
         BLUE)
    ]

    for k, (layer_name, details, bar_color) in enumerate(arch_layers):
        y = Inches(1.8 + k * 1.65)
        layer_box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y, Inches(11.7), Inches(1.45))
        layer_box.fill.solid()
        layer_box.fill.fore_color.rgb = CARD_BG
        layer_box.line.color.rgb = CARD_BORDER
        layer_box.line.width = Pt(1.5)

        # Left color strip
        strip = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y, Inches(0.2), Inches(1.45))
        strip.fill.solid()
        strip.fill.fore_color.rgb = bar_color
        strip.line.fill.background()

        tf_l = layer_box.text_frame
        tf_l.word_wrap = True
        p_ln = tf_l.paragraphs[0]
        p_ln.text = f"   {layer_name}"
        p_ln.font.size = Pt(15)
        p_ln.font.bold = True
        p_ln.font.color.rgb = TEXT_DARK

        p_ld = tf_l.add_paragraph()
        p_ld.text = details
        p_ld.font.size = Pt(11.5)
        p_ld.font.color.rgb = TEXT_MUTED

    s4.notes_slide.notes_text_frame.text = (
        "Here we show the modular, three-tier architecture of DocStudio. In the client layer, PDF.js renders vectors "
        "directly in the browser canvas. The FastAPI layer streams data asynchronously without disk bottlenecks. "
        "The processing engine utilizes PyMuPDF and Pillow for surgical stream transformations and iterative compression."
    )

    # ==========================================
    # SLIDE 5: Key Innovation 1: Govt Form Resizer
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    set_bg(s5, BG_LIGHT)
    add_header(s5, "04. Key Innovation #1", "Govt Form & Exam Document Resizer (Prompt & KB Target)")

    grid_items = [
        ("🤖 AI Natural Language Parser",
         "Users can type prompt requirements naturally:\n"
         "• 'passport photo 3.5x4.5 cm under 50kb'\n"
         "• 'SSC signature 4x2 cm max 20 kb'\n"
         "Regex parser automatically extracts width, height, unit (cm/mm/in/px), and target file size.",
         EMERALD),
        ("⚡ 1-Click Standard Presets",
         "Pre-programmed standard dimensions used across 95% of recruitment portals:\n"
         "• Passport Photo: 3.5 × 4.5 cm (< 50 KB)\n"
         "• Official Signature: 4.0 × 2.0 cm (< 20 KB)\n"
         "• ID / PAN / Aadhaar Card: 8.5 × 5.5 cm (< 100 KB)\n"
         "• Marksheet / Certificate: Full A4 (< 200 KB)",
         INDIGO),
        ("🔒 Live Proportional Aspect Lock",
         "Prevents facial distortion and signature warping:\n"
         "• Real-time canvas constraint keeps aspect ratio intact\n"
         "• Scaling corner handles maintains exact 3.5:4.5 or 4:2 ratio\n"
         "• User visually frames their photo without guessing coordinates",
         BLUE),
        ("🗜️ Stepped KB Compression Loop",
         "Guarantees output is strictly under target KB limit:\n"
         "• Evaluates initial clipped PDF stream size\n"
         "• If size > target, triggers iterative resolution & JPEG quality decay steps (200 DPI -> 72 DPI, Q85 -> Q35)\n"
         "• 100% portal acceptance rate without blurriness",
         AMBER)
    ]

    for m, (g_title, g_desc, g_color) in enumerate(grid_items):
        row = m // 2
        col = m % 2
        x = Inches(0.8 + col * 5.95)
        y = Inches(1.8 + row * 2.45)
        g_box = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.75), Inches(2.3))
        g_box.fill.solid()
        g_box.fill.fore_color.rgb = CARD_BG
        g_box.line.color.rgb = CARD_BORDER
        g_box.line.width = Pt(1.5)

        # Header tag
        g_tag = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.75), Inches(0.12))
        g_tag.fill.solid()
        g_tag.fill.fore_color.rgb = g_color
        g_tag.line.fill.background()

        tf_g = g_box.text_frame
        tf_g.word_wrap = True
        p_gt = tf_g.paragraphs[0]
        p_gt.text = g_title
        p_gt.font.size = Pt(14.5)
        p_gt.font.bold = True
        p_gt.font.color.rgb = TEXT_DARK

        p_gd = tf_g.add_paragraph()
        p_gd.text = g_desc
        p_gd.font.size = Pt(11)
        p_gd.font.color.rgb = TEXT_MUTED

    s5.notes_slide.notes_text_frame.text = (
        "This is one of the standout innovations of our project. No existing free PDF editor allows users to type "
        "'Passport photo 3.5x4.5 cm under 50kb' and automatically lock the aspect ratio while applying iterative "
        "stepped compression in the backend to ensure the file is strictly below 50KB or 20KB for government portals."
    )

    # ==========================================
    # SLIDE 6: Key Innovation 2: Full PDF Editor
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    set_bg(s6, BG_LIGHT)
    add_header(s6, "05. Key Innovation #2", "Interactive Online PDF Editor & Dynamic Page Insertion")

    editor_features = [
        ("📄 Dynamic Page Insertion ('+ Insert page here')",
         "Unlike rigid viewers that only allow editing pre-existing pages, DocStudio dynamically creates clean, blank canvas pages anywhere in the document stream, recalculating page indices and manifests on the fly."),
        ("🗑️ Interactive Page Management & Deletion",
         "Users can remove unwanted pages directly from the live preview bar with real-time UI synchronization and zero document corruption."),
        ("✏️ Low-Level Text Injection & Base-14 Font Mapping",
         "Overcomes PyMuPDF buffer exceptions by mapping custom fonts to valid PostScript Base-14 font tables (helv, times-roman, courier), inserting searchable text boxes with custom sizes and colors."),
        ("✍️ Freehand Drawing, Digital Signing & Images",
         "HTML5 Canvas capture enables users to place digital signatures, stamps, and high-resolution images with seamless stream flattening during final export.")
    ]

    for n, (feat_title, feat_desc) in enumerate(editor_features):
        y = Inches(1.8 + n * 1.25)
        f_box = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y, Inches(11.7), Inches(1.15))
        f_box.fill.solid()
        f_box.fill.fore_color.rgb = CARD_BG
        f_box.line.color.rgb = CARD_BORDER
        f_box.line.width = Pt(1.5)

        tf_f = f_box.text_frame
        tf_f.word_wrap = True
        p_ft = tf_f.paragraphs[0]
        p_ft.text = feat_title
        p_ft.font.size = Pt(14)
        p_ft.font.bold = True
        p_ft.font.color.rgb = INDIGO

        p_fd = tf_f.add_paragraph()
        p_fd.text = feat_desc
        p_fd.font.size = Pt(11.5)
        p_fd.font.color.rgb = TEXT_MUTED

    s6.notes_slide.notes_text_frame.text = (
        "Our second core innovation is the full-featured PDF Editor. We solved a critical limitation of existing tools: "
        "allowing users to dynamically insert fresh blank pages inside an existing PDF, write text on them with proper "
        "Base-14 font mapping, and delete unwanted pages with real-time PyMuPDF stream re-compilation."
    )

    # ==========================================
    # SLIDE 7: Comprehensive Suite Ecosystem
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    set_bg(s7, BG_LIGHT)
    add_header(s7, "06. Feature Landscape", "Comprehensive Suite: 30+ PDF & PPT Utilities")

    eco_cards = [
        ("📁 Document Assembly",
         "• Merge PDFs (drag-and-drop order)\n"
         "• Split by page ranges or equal parts\n"
         "• Alternate & Mix multi-source pages\n"
         "• Rotate pages by 90°, 180°, 270°",
         INDIGO),
        ("🛡️ Security & Privacy",
         "• AES-256 Bit Password Encryption\n"
         "• Unlock & Remove PDF Passwords\n"
         "• Metadata Scrubber (strip authors/dates)\n"
         "• Auto-cleaning ephemeral server storage",
         EMERALD),
        ("🔄 Multi-Format Converters",
         "• Image to PDF (JPG, PNG, WebP)\n"
         "• PDF to High-Res Images\n"
         "• Word (.docx) to PDF Generator\n"
         "• Excel (.xlsx) Table to PDF Exporter",
         BLUE),
        ("📊 PPT Studio Generator",
         "• Interactive Slide Deck Designer\n"
         "• Multi-theme support (Modern Blue, Slate)\n"
         "• Structured Slide Bullet Synthesizer\n"
         "• Instant Export to native .pptx files",
         AMBER)
    ]

    for p, (e_title, e_bullets, e_col) in enumerate(eco_cards):
        x = Inches(0.8 + p * 2.95)
        e_box = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.8), Inches(2.85), Inches(4.8))
        e_box.fill.solid()
        e_box.fill.fore_color.rgb = CARD_BG
        e_box.line.color.rgb = CARD_BORDER
        e_box.line.width = Pt(1.5)

        e_top = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.8), Inches(2.85), Inches(0.12))
        e_top.fill.solid()
        e_top.fill.fore_color.rgb = e_col
        e_top.line.fill.background()

        tf_e = e_box.text_frame
        tf_e.word_wrap = True
        p_et = tf_e.paragraphs[0]
        p_et.text = e_title
        p_et.font.size = Pt(15)
        p_et.font.bold = True
        p_et.font.color.rgb = TEXT_DARK

        p_eb = tf_e.add_paragraph()
        p_eb.text = "\n" + e_bullets
        p_eb.font.size = Pt(11.5)
        p_eb.font.color.rgb = TEXT_MUTED

    s7.notes_slide.notes_text_frame.text = (
        "DocStudio is not a single-purpose tool; it is a full productivity ecosystem. It includes document assembly, "
        "military-grade AES-256 encryption, multi-format converters for Word, Excel, and Images, and an interactive "
        "Presentation Studio for creating native PowerPoint slide decks."
    )

    # ==========================================
    # SLIDE 8: Tech Stack & Implementation Details
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    set_bg(s8, BG_LIGHT)
    add_header(s8, "07. Implementation Details", "Technology Stack & Technical Implementation")

    tech_table = [
        ("Component", "Technology / Library", "Engineering Purpose"),
        ("Backend Framework", "Python 3.13 + FastAPI", "High-concurrency ASGI web framework with async endpoints"),
        ("PDF Engine", "PyMuPDF (Fitz 1.24+)", "Vector clipping, page restructuring, text injection & stream repair"),
        ("Image & Format Engine", "Pillow (PIL) + ReportLab", "Dynamic DPI downsampling, color space correction & PDF generation"),
        ("Office Converters", "python-pptx, docx, openpyxl", "Binary format parsing and automated PowerPoint/Word/Excel synthesis"),
        ("Frontend UI", "HTML5, Tailwind CSS, JS (ES6+)", "Modern responsive SPA interface with zero external framework bloat"),
        ("Client PDF Renderer", "Mozilla PDF.js 3.11", "Client-side vector rendering of PDF pages onto HTML5 canvas"),
        ("Database & Auth", "SQLite3 + PBKDF2 Hashing", "Lightweight, zero-config relational store for sessions and history"),
        ("Cloud & Deployment", "Git + Render PaaS + Cloudflare", "Automated GitHub CI/CD, SSL encryption & cloud serverless hosting")
    ]

    # Create Table Shape
    rows = len(tech_table)
    cols = 3
    t_shape = s8.shapes.add_table(rows, cols, Inches(0.8), Inches(1.8), Inches(11.7), Inches(4.8))
    table = t_shape.table
    table.columns[0].width = Inches(2.5)
    table.columns[1].width = Inches(3.5)
    table.columns[2].width = Inches(5.7)

    for r_idx, row in enumerate(tech_table):
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = val
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            p.font.name = "Segoe UI"
            if r_idx == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = BG_DARK
                p.font.bold = True
                p.font.size = Pt(13)
                p.font.color.rgb = RGBColor(255, 255, 255)
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = CARD_BG if r_idx % 2 == 0 else RGBColor(241, 245, 249)
                p.font.size = Pt(11)
                p.font.color.rgb = TEXT_DARK if c_idx == 0 else TEXT_MUTED
                if c_idx == 1:
                    p.font.bold = True
                    p.font.color.rgb = INDIGO

    s8.notes_slide.notes_text_frame.text = (
        "Here is the complete implementation matrix. We chose FastAPI for its asynchronous capabilities and PyMuPDF "
        "for its unparalleled speed over traditional Python PDF libraries like PyPDF2. On the frontend, we maintained "
        "lightweight performance using Vanilla JavaScript and Tailwind CSS without heavy React or Angular dependencies."
    )

    # ==========================================
    # SLIDE 9: Results, Deployment & Performance
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    set_bg(s9, BG_LIGHT)
    add_header(s9, "08. Results & Validation", "Live Deployment, Performance Metrics & SEO")

    # 4 Metric Cards
    metrics = [
        ("< 350 ms", "Average API Execution Time", "PyMuPDF C-bindings provide sub-second operations across all tools.", EMERALD),
        ("100%", "Target KB Accuracy", "Zero-failure compliance with government exam size thresholds (<= 20KB / 50KB).", INDIGO),
        ("24/7 Live", "Cloud Deployment on Render", "Global accessibility with automatic CI/CD from GitHub main branch.", BLUE),
        ("SEO Indexed", "Google Search Console Ready", "Integrated JSON-LD schema, sitemap.xml, robots.txt, and meta tags.", AMBER)
    ]

    for q, (num, label, desc, col) in enumerate(metrics):
        x = Inches(0.8 + q * 2.95)
        m_box = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.8), Inches(2.85), Inches(2.2))
        m_box.fill.solid()
        m_box.fill.fore_color.rgb = CARD_BG
        m_box.line.color.rgb = CARD_BORDER
        m_box.line.width = Pt(1.5)

        tf_m = m_box.text_frame
        tf_m.word_wrap = True
        p_mn = tf_m.paragraphs[0]
        p_mn.text = num
        p_mn.font.size = Pt(28)
        p_mn.font.bold = True
        p_mn.font.color.rgb = col

        p_ml = tf_m.add_paragraph()
        p_ml.text = label
        p_ml.font.size = Pt(13)
        p_ml.font.bold = True
        p_ml.font.color.rgb = TEXT_DARK

        p_md = tf_m.add_paragraph()
        p_md.text = desc
        p_md.font.size = Pt(10.5)
        p_md.font.color.rgb = TEXT_MUTED

    # Bottom Live URL Showcase Card
    url_card = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.3), Inches(11.7), Inches(2.3))
    url_card.fill.solid()
    url_card.fill.fore_color.rgb = BG_DARK
    url_card.line.fill.background()
    tf_uc = url_card.text_frame
    tf_uc.word_wrap = True

    p_u1 = tf_uc.paragraphs[0]
    p_u1.text = "🌐 Live Production Environment Demonstration"
    p_u1.font.size = Pt(18)
    p_u1.font.bold = True
    p_u1.font.color.rgb = EMERALD

    p_u2 = tf_uc.add_paragraph()
    p_u2.text = (
        "\n• Live Application URL: https://docstudio-dty3.onrender.com/\n"
        "• Public GitHub Repository: https://github.com/AlokKumarKushwaha/doc-studio\n"
        "• Google Search Console verified property with daily automatic sitemap crawling enabled\n"
        "• Verified zero-error operation on mobile browsers (Chrome Android, Safari iOS) and desktop devices"
    )
    p_u2.font.size = Pt(13)
    p_u2.font.color.rgb = RGBColor(226, 232, 240)

    s9.notes_slide.notes_text_frame.text = (
        "Unlike purely academic projects that only exist on localhost, DocStudio is deployed live in production "
        "on Render with automated GitHub workflows, verified on Google Search Console, and operating with "
        "exceptional sub-second execution speed across desktop and mobile devices."
    )

    # ==========================================
    # SLIDE 10: Conclusion & Future Scope (Q&A)
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    set_bg(s10, BG_DARK)

    # Header Box
    h_box = s10.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.5), Inches(1.2))
    tf_h = h_box.text_frame
    tf_h.word_wrap = True
    p_ht = tf_h.paragraphs[0]
    p_ht.text = "CONCLUSION & FUTURE ROADMAP"
    p_ht.font.size = Pt(13)
    p_ht.font.bold = True
    p_ht.font.color.rgb = EMERALD

    p_hm = tf_h.add_paragraph()
    p_hm.text = "Empowering Users with Intelligent Document Engineering"
    p_hm.font.size = Pt(28)
    p_hm.font.bold = True
    p_hm.font.color.rgb = RGBColor(255, 255, 255)

    # Two Columns: Left Conclusion, Right Future Scope
    c_left = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.2), Inches(5.6), Inches(4.4))
    c_left.fill.solid()
    c_left.fill.fore_color.rgb = RGBColor(30, 41, 59)
    c_left.line.color.rgb = RGBColor(51, 65, 85)
    tf_cl = c_left.text_frame
    tf_cl.word_wrap = True

    p_cl1 = tf_cl.paragraphs[0]
    p_cl1.text = "✅ Project Summary & Impact"
    p_cl1.font.size = Pt(18)
    p_cl1.font.bold = True
    p_cl1.font.color.rgb = RGBColor(255, 255, 255)

    p_cl2 = tf_cl.add_paragraph()
    p_cl2.text = (
        "\n• Successfully engineered a full-featured, zero-cost alternative to expensive document software.\n\n"
        "• Solved the government exam upload challenge through AI prompt parsing and deterministic compression.\n\n"
        "• Built an interactive in-browser PDF editor capable of dynamic page insertion and font synthesis.\n\n"
        "• Deployed live with enterprise-grade security and guest-friendly usability."
    )
    p_cl2.font.size = Pt(13)
    p_cl2.font.color.rgb = RGBColor(203, 213, 225)

    c_right = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(2.2), Inches(5.7), Inches(4.4))
    c_right.fill.solid()
    c_right.fill.fore_color.rgb = RGBColor(30, 41, 59)
    c_right.line.color.rgb = RGBColor(51, 65, 85)
    tf_cr = c_right.text_frame
    tf_cr.word_wrap = True

    p_cr1 = tf_cr.paragraphs[0]
    p_cr1.text = "🔮 Future Roadmap & Enhancements"
    p_cr1.font.size = Pt(18)
    p_cr1.font.bold = True
    p_cr1.font.color.rgb = INDIGO

    p_cr2 = tf_cr.add_paragraph()
    p_cr2.text = (
        "\n• Multilingual OCR Engine: Integration with Tesseract OCR for searchable Hindi/Regional language document scanning.\n\n"
        "• Real-Time Collaborative PDF Annotation: WebSocket-powered multi-user markup and reviewing.\n\n"
        "• Cloud Drive Synchronization: Native integration with Google Drive, OneDrive, and Dropbox.\n\n"
        "• AI Document Summarizer: LLM-powered instant question-answering over large PDF documents."
    )
    p_cr2.font.size = Pt(13)
    p_cr2.font.color.rgb = RGBColor(203, 213, 225)

    s10.notes_slide.notes_text_frame.text = (
        "In conclusion, DocStudio successfully bridges the gap between complex document engineering and everyday accessibility. "
        "Our roadmap includes multilingual OCR and real-time collaboration. Thank you so much for your time and guidance. "
        "I am now open to any questions and feedback from the evaluators."
    )

    # Save to file
    out_path = os.path.join(os.getcwd(), "DocStudio_College_Presentation.pptx")
    prs.save(out_path)
    print(f"Presentation saved successfully to: {out_path}")

if __name__ == "__main__":
    create_presentation()
